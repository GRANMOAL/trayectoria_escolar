import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .middleware import PRIVACY_NOTICE_VERSION
from .privacidad import alumnos_con_consentimiento
from .models import (
    Alumno, Asignatura, AvisoPrivacidadAceptacion, Calificacion, Consentimiento,
    Grupo, Nota, Semestre, SolicitudARCO, SolicitudARCOEvento,
    PronosticoLog, SolicitudRevisionPronostico,
)
from .pronosticos import explicar_pronostico, metricas_validacion


class AuthenticationTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='docente',
            password='clave-segura-prueba',
        )

    def test_application_redirects_anonymous_users_to_login(self):
        response = self.client.get(reverse('dashboard'))

        self.assertRedirects(
            response,
            f"{reverse('login')}?next={reverse('dashboard')}",
        )

    def test_application_api_requires_login(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse('agente_estado_api'))

        self.assertRedirects(
            response,
            f"{reverse('privacy_acceptance')}?next={reverse('agente_estado_api')}",
        )

    def test_login_requires_privacy_acceptance_before_dashboard(self):
        response = self.client.post(reverse('login'), {
            'username': self.user.username,
            'password': 'clave-segura-prueba',
        })

        self.assertRedirects(
            response,
            f"{reverse('privacy_acceptance')}?next={reverse('dashboard')}",
        )
        self.assertEqual(self.client.get(reverse('dashboard')).status_code, 302)
        self.assertEqual(AvisoPrivacidadAceptacion.objects.count(), 0)

    def test_accepting_privacy_notice_records_user_and_version(self):
        self.client.force_login(self.user)

        response = self.client.post(reverse('privacy_acceptance'), {
            'decision': 'aceptar',
            'aceptacion': 'on',
        })

        self.assertRedirects(response, reverse('dashboard'))
        acceptance = AvisoPrivacidadAceptacion.objects.get()
        self.assertEqual(acceptance.usuario, self.user)
        self.assertEqual(acceptance.version, PRIVACY_NOTICE_VERSION)
        self.assertIsNotNone(acceptance.aceptado_en)
        self.assertEqual(self.client.get(reverse('dashboard')).status_code, 200)

    def test_rejecting_privacy_notice_logs_user_out(self):
        self.client.force_login(self.user)

        response = self.client.post(reverse('privacy_acceptance'), {
            'decision': 'rechazar',
        })

        self.assertRedirects(response, reverse('login'))
        self.assertEqual(AvisoPrivacidadAceptacion.objects.count(), 0)
        self.assertRedirects(
            self.client.get(reverse('dashboard')),
            f"{reverse('login')}?next={reverse('dashboard')}",
        )

    def test_acceptance_requires_confirmation_checkbox(self):
        self.client.force_login(self.user)

        response = self.client.post(reverse('privacy_acceptance'), {
            'decision': 'aceptar',
        })

        self.assertRedirects(response, reverse('login'))
        self.assertEqual(AvisoPrivacidadAceptacion.objects.count(), 0)

    def test_logout_ends_the_session(self):
        self.client.force_login(self.user)

        response = self.client.post(reverse('logout'))

        self.assertRedirects(response, reverse('login'))
        self.assertRedirects(
            self.client.get(reverse('dashboard')),
            f"{reverse('login')}?next={reverse('dashboard')}",
        )

    def test_privacy_notice_is_public(self):
        response = self.client.get(reverse('aviso_privacidad'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Aviso de privacidad')
        self.assertContains(response, 'Aviso de privacidad simplificado')
        self.assertContains(response, 'Aviso de privacidad integral')
        self.assertContains(response, PRIVACY_NOTICE_VERSION)
        self.assertContains(response, '26 de octubre de 2026')

    def test_login_footer_links_to_privacy_notice(self):
        response = self.client.get(reverse('login'))

        self.assertContains(response, 'Aviso de privacidad · versión integral y simplificada')


class StudentConsentTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='docente-consentimientos',
            password='clave-segura-prueba',
        )
        self.alumno = Alumno.objects.create(
            numero_cuenta='2026001',
            nombre_completo='Alumno Sin Consentimiento',
        )

    def iniciar_sesion_aceptada(self):
        self.client.force_login(self.user)
        AvisoPrivacidadAceptacion.objects.create(
            usuario=self.user,
            version=PRIVACY_NOTICE_VERSION,
        )

    def registrar_consentimiento(self):
        return self.client.post(reverse('gestionar_consentimientos'), {
            'alumno': self.alumno.id,
            'tutor_nombre': 'Madre de prueba',
            'parentesco': 'madre',
            'medio': 'presencial',
            'aceptacion': 'on',
        })

    def test_consent_management_requires_login(self):
        response = self.client.get(reverse('gestionar_consentimientos'))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse('login'), response['Location'])

    def test_registering_guardian_consent_enables_student_data(self):
        self.iniciar_sesion_aceptada()
        lista_url = reverse('lista_alumnos')
        self.assertContains(
            self.client.get(reverse('gestionar_consentimientos')),
            'Registrar consentimiento',
        )
        self.assertNotContains(self.client.get(lista_url), self.alumno.nombre_completo)

        response = self.registrar_consentimiento()

        self.assertRedirects(response, reverse('gestionar_consentimientos'))
        consentimiento = Consentimiento.objects.get(alumno=self.alumno)
        self.assertEqual(consentimiento.tutor_nombre, 'Madre de prueba')
        self.assertEqual(consentimiento.parentesco, 'madre')
        self.assertEqual(consentimiento.version_aviso, PRIVACY_NOTICE_VERSION)
        self.assertEqual(consentimiento.medio, 'presencial')
        self.assertEqual(consentimiento.estado, Consentimiento.ESTADO_VIGENTE)
        self.assertIsNotNone(consentimiento.fecha_hora)
        self.assertEqual(consentimiento.registrado_por, self.user)
        self.assertContains(self.client.get(lista_url), self.alumno.nombre_completo)

    def test_revoking_consent_blocks_student_data_again(self):
        self.iniciar_sesion_aceptada()
        self.registrar_consentimiento()
        consentimiento = Consentimiento.objects.get(alumno=self.alumno)

        response = self.client.post(
            reverse('revocar_consentimiento', args=[consentimiento.id]),
        )

        self.assertRedirects(response, reverse('gestionar_consentimientos'))
        consentimiento.refresh_from_db()
        self.assertEqual(consentimiento.estado, Consentimiento.ESTADO_REVOCADO)
        self.assertIsNotNone(consentimiento.revocado_en)
        self.assertNotContains(
            self.client.get(reverse('lista_alumnos')),
            self.alumno.nombre_completo,
        )

    def test_direct_student_api_hides_students_without_consent(self):
        self.iniciar_sesion_aceptada()

        response = self.client.get(reverse('seguimiento_notas_api'), {
            'alumno_id': self.alumno.id,
        })

        self.assertEqual(response.status_code, 404)

    def test_ai_api_rejects_unconsented_student_selection(self):
        self.iniciar_sesion_aceptada()

        response = self.client.post(
            reverse('agente_chat_api'),
            data=json.dumps({
                'mensajes': [{'role': 'user', 'content': 'Resume el avance'}],
                'alumno_id': self.alumno.id,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 403)

    @override_settings(REQUIRE_STUDENT_CONSENT=False)
    def test_student_data_block_can_be_disabled(self):
        self.iniciar_sesion_aceptada()

        response = self.client.get(reverse('lista_alumnos'))

        self.assertContains(response, self.alumno.nombre_completo)

    def test_admin_hides_unconsented_students_and_shows_privacy_footer(self):
        admin_user = get_user_model().objects.create_superuser(
            username='admin-consentimientos',
            email='admin@example.test',
            password='clave-segura-prueba',
        )
        self.client.force_login(admin_user)
        AvisoPrivacidadAceptacion.objects.create(
            usuario=admin_user,
            version=PRIVACY_NOTICE_VERSION,
        )

        response = self.client.get(reverse('admin:escolar_alumno_changelist'))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, self.alumno.nombre_completo)
        self.assertContains(response, 'Aviso de privacidad')


class GradeRangeTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='docente-calificaciones',
            password='clave-segura-prueba',
        )
        self.semestre = Semestre.objects.create(
            nombre='2026-B',
            anio=2026,
            periodo='AGO-DIC',
        )
        self.grupo = Grupo.objects.create(nombre='B', semestre=self.semestre)
        self.alumno = Alumno.objects.create(
            numero_cuenta='2026003',
            nombre_completo='Alumno Calificaciones',
            grupo=self.grupo,
        )
        self.asignatura = Asignatura.objects.create(
            nombre='Ciencias',
            clave='CIE-01',
            semestre=self.semestre,
        )
        Consentimiento.objects.create(
            alumno=self.alumno,
            tutor_nombre='Tutor de prueba',
            parentesco='tutor',
            version_aviso=PRIVACY_NOTICE_VERSION,
            medio='presencial',
        )
        self.client.force_login(self.user)
        AvisoPrivacidadAceptacion.objects.create(
            usuario=self.user,
            version=PRIVACY_NOTICE_VERSION,
        )

    def guardar(self, valores):
        return self.client.post(
            reverse('guardar_calificaciones'),
            data=json.dumps({
                'asignatura_id': self.asignatura.id,
                'parcial': 1,
                'rows': [{
                    'alumno_id': self.alumno.id,
                    'hetero': valores[0],
                    'co': valores[1],
                    'auto': valores[2],
                }],
            }),
            content_type='application/json',
        )

    def test_saves_decimal_scores_within_range(self):
        response = self.guardar(['8.25', '9.5', '0.01'])

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['ok'])
        calificacion = Calificacion.objects.get(alumno=self.alumno)
        self.assertEqual(calificacion.heteroevaluacion, Decimal('8.25'))
        self.assertEqual(calificacion.coevaluacion, Decimal('9.5'))
        self.assertEqual(calificacion.autoevaluacion, Decimal('0.01'))

    def test_rejects_out_of_range_scores_without_partial_saves(self):
        for valor in ['-0.01', '10.01']:
            with self.subTest(valor=valor):
                response = self.guardar([valor, '5', '5'])

                self.assertEqual(response.status_code, 400)
                self.assertFalse(response.json()['ok'])
                self.assertIn('entre 0 y 10', response.json()['error'])
                self.assertFalse(Calificacion.objects.exists())

    def test_rejects_non_finite_or_excess_precision_scores(self):
        for valor in ['NaN', 'Infinity', '8.123']:
            with self.subTest(valor=valor):
                response = self.guardar([valor, '7', '8'])

                self.assertEqual(response.status_code, 400)
                self.assertFalse(Calificacion.objects.exists())

    def test_capture_page_displays_range_and_decimal_limits(self):
        response = self.client.get(reverse('captura_rapida'), {
            'grupo': self.grupo.id,
            'asignatura': self.asignatura.id,
        })

        self.assertContains(response, 'Calificaciones permitidas: de 0 a 10.')
        self.assertContains(response, 'Se admiten hasta dos decimales.')
        self.assertContains(response, 'step="0.01"', count=3)


class ResponsibleForecastTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='coordinacion-pronosticos',
            password='clave-segura-prueba',
            is_staff=True,
        )
        self.semestre = Semestre.objects.create(
            nombre='2026-C',
            anio=2026,
            periodo='ENE-JUN',
        )
        self.grupo = Grupo.objects.create(nombre='C', semestre=self.semestre)
        self.alumno = Alumno.objects.create(
            numero_cuenta='2026004',
            nombre_completo='Alumno Pronóstico',
            email='alumno@example.test',
            grupo=self.grupo,
        )
        self.asignatura = Asignatura.objects.create(
            nombre='Historia',
            clave='HIS-01',
            semestre=self.semestre,
        )
        Consentimiento.objects.create(
            alumno=self.alumno,
            tutor_nombre='Tutor de prueba',
            parentesco='tutor',
            version_aviso=PRIVACY_NOTICE_VERSION,
            medio='presencial',
        )
        self.calificaciones = [
            Calificacion.objects.create(
                alumno=self.alumno,
                asignatura=self.asignatura,
                parcial=parcial,
                heteroevaluacion=promedio,
                coevaluacion=promedio,
                autoevaluacion=promedio,
            )
            for parcial, promedio in ((1, 6), (2, 8), (3, 9))
        ]
        self.client.force_login(self.user)
        AvisoPrivacidadAceptacion.objects.create(
            usuario=self.user,
            version=PRIVACY_NOTICE_VERSION,
        )

    def generar_pronostico(self):
        return self.client.get(
            reverse('pronostico'),
            {'grupo': self.grupo.id},
        )

    def test_forecast_shows_caveat_explanation_metrics_and_does_not_change_grades(self):
        antes = list(
            Calificacion.objects.filter(alumno=self.alumno)
            .order_by('parcial')
            .values_list('parcial', 'heteroevaluacion', 'coevaluacion', 'autoevaluacion')
        )

        response = self.generar_pronostico()

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Estimación orientativa; requiere revisión humana.')
        self.assertContains(response, 'Factores y coeficientes')
        self.assertContains(response, 'Coef. × valor')
        self.assertContains(response, 'R²:')
        self.assertContains(response, 'MAE:')
        self.assertContains(response, f'?pronostico={PronosticoLog.objects.get(alumno=self.alumno).id}')
        log = PronosticoLog.objects.get(alumno=self.alumno, asignatura=self.asignatura)
        self.assertEqual(log.salida['pronostico_p3'], 10)
        self.assertEqual(log.salida['factores'][0]['coeficiente'], -1)
        self.assertEqual(log.salida['factores'][1]['coeficiente'], 2)
        despues = list(
            Calificacion.objects.filter(alumno=self.alumno)
            .order_by('parcial')
            .values_list('parcial', 'heteroevaluacion', 'coevaluacion', 'autoevaluacion')
        )
        self.assertEqual(despues, antes)

    def test_forecast_log_contains_only_permitted_aggregated_features(self):
        self.generar_pronostico()
        log = PronosticoLog.objects.get(alumno=self.alumno)

        self.assertEqual(log.version_modelo, 'tendencia-lineal-v1')
        self.assertEqual(
            set(log.entradas_agregadas),
            {'parciales'},
        )
        self.assertEqual(
            set(log.entradas_agregadas['parciales'][0]),
            {'parcial', 'promedio'},
        )
        for atributo_protegido in (
            'genero', 'edad', 'origen', 'discapacidad', 'nivel_socioeconomico',
        ):
            self.assertNotIn(atributo_protegido, json.dumps(log.entradas_agregadas).lower())

    def test_forecast_excel_exports_use_same_estimate_and_responsible_caveat(self):
        response = self.client.get(
            reverse('exportar_pronostico'),
            {'grupo': self.grupo.id},
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook['Pronóstico General']

        self.assertEqual(response.status_code, 200)
        self.assertIn('Estimación orientativa; requiere revisión humana', worksheet['A2'].value)
        self.assertEqual(worksheet['F4'].value, 10)
        self.assertIn('Promedio P1: -1', worksheet['H4'].value)

        index_response = self.client.get(
            reverse('exportar_indices_pronostico'),
            {'grupo': self.grupo.id},
        )
        index_workbook = load_workbook(BytesIO(index_response.content), data_only=True)
        index_worksheet = index_workbook['Índice Pronóstico']
        self.assertIn('Estimación orientativa; requiere revisión humana', index_worksheet['A2'].value)
        self.assertEqual(index_worksheet['D4'].value, 1)
        self.assertEqual(index_worksheet['E4'].value, 0)
        self.assertEqual(PronosticoLog.objects.filter(alumno=self.alumno).count(), 1)

    def test_student_or_guardian_can_submit_human_review_request(self):
        self.generar_pronostico()
        log = PronosticoLog.objects.get(alumno=self.alumno)
        self.client.logout()
        url = reverse('solicitar_revision_pronostico')

        page = self.client.get(url, {'pronostico': log.id})
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'Solicitar revisión humana')

        response = self.client.post(url, {
            'pronostico': log.id,
            'numero_cuenta': self.alumno.numero_cuenta,
            'solicitante': 'Madre de prueba',
            'parentesco': 'madre',
            'contacto_email': 'madre@example.test',
            'motivo': 'Solicito revisar los factores del cálculo.',
            'confirmacion': 'on',
        })

        solicitud = SolicitudRevisionPronostico.objects.get()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'REV-{solicitud.id}')
        self.assertEqual(solicitud.alumno, self.alumno)
        self.assertEqual(solicitud.estado, 'recibida')
        self.assertEqual(solicitud.eventos.count(), 1)

    def test_public_review_entry_is_available_from_login_without_forecast_id(self):
        self.generar_pronostico()
        self.client.logout()
        url = reverse('solicitar_revision_pronostico')

        self.assertContains(
            self.client.get(reverse('login')),
            'Solicitar revisión humana de pronóstico',
        )
        self.assertEqual(self.client.get(url).status_code, 200)
        response = self.client.post(url, {
            'numero_cuenta': self.alumno.numero_cuenta,
            'solicitante': 'Alumno de prueba',
            'parentesco': 'alumno',
            'contacto_email': 'alumno@example.test',
            'motivo': 'Solicito revisar mi pronóstico.',
            'confirmacion': 'on',
        })

        solicitud = SolicitudRevisionPronostico.objects.get()
        self.assertContains(response, f'REV-{solicitud.id}')
        self.assertEqual(solicitud.pronostico.alumno, self.alumno)

    def test_review_request_requires_valid_student_number(self):
        self.generar_pronostico()
        log = PronosticoLog.objects.get(alumno=self.alumno)
        self.client.logout()

        response = self.client.post(reverse('solicitar_revision_pronostico'), {
            'pronostico': log.id,
            'numero_cuenta': 'incorrecto',
            'solicitante': 'Madre de prueba',
            'parentesco': 'madre',
            'contacto_email': 'madre@example.test',
            'motivo': 'Solicito revisión.',
            'confirmacion': 'on',
        })

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'No fue posible validar el pronóstico')
        self.assertFalse(SolicitudRevisionPronostico.objects.exists())

    def test_coordination_tracks_review_status_and_response(self):
        self.generar_pronostico()
        log = PronosticoLog.objects.get(alumno=self.alumno)
        solicitud = SolicitudRevisionPronostico.objects.create(
            alumno=self.alumno,
            pronostico=log,
            solicitante='Tutor de prueba',
            parentesco='tutor',
            contacto_email='tutor@example.test',
            motivo='Solicito revisión.',
        )

        response = self.client.post(reverse('coordinacion_pronosticos'), {
            'solicitud': solicitud.id,
            'estado': 'atendida',
            'respuesta': 'Se revisaron los datos y la estimación.',
        })

        self.assertRedirects(
            response,
            f"{reverse('coordinacion_pronosticos')}?solicitud={solicitud.id}",
        )
        solicitud.refresh_from_db()
        self.assertEqual(solicitud.estado, 'atendida')
        self.assertEqual(solicitud.responsable, self.user)
        self.assertTrue(solicitud.fecha_actualizacion)
        self.assertEqual(solicitud.eventos.count(), 1)

    def test_forecast_review_panel_is_limited_to_staff(self):
        regular_user = get_user_model().objects.create_user(
            username='docente-sin-coordinacion',
            password='clave-segura-prueba',
        )
        self.client.force_login(regular_user)
        AvisoPrivacidadAceptacion.objects.create(
            usuario=regular_user,
            version=PRIVACY_NOTICE_VERSION,
        )

        response = self.client.get(reverse('coordinacion_pronosticos'))

        self.assertEqual(response.status_code, 403)

    def test_forecast_helpers_report_metrics_and_require_both_partials(self):
        self.assertIsNone(explicar_pronostico({1: 7}))
        self.assertEqual(metricas_validacion([]), {
            'r2': None,
            'mae': None,
            'muestras': 0,
        })
        self.assertEqual(
            metricas_validacion([(8, 9), (6, 7)]),
            {'r2': 0.0, 'mae': 1.0, 'muestras': 2},
        )


class ArcoTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='coordinacion-arco',
            password='clave-segura-prueba',
            is_staff=True,
        )
        self.semestre = Semestre.objects.create(
            nombre='2026-A',
            anio=2026,
            periodo='ENE-JUN',
        )
        self.grupo = Grupo.objects.create(nombre='A', semestre=self.semestre)
        self.alumno = Alumno.objects.create(
            numero_cuenta='2026002',
            nombre_completo='Titular ARCO',
            email='titular@example.test',
            grupo=self.grupo,
        )
        self.asignatura = Asignatura.objects.create(
            nombre='Matemáticas',
            clave='MAT-ARCO',
            semestre=self.semestre,
        )
        self.calificacion = Calificacion.objects.create(
            alumno=self.alumno,
            asignatura=self.asignatura,
            parcial=1,
            heteroevaluacion=8,
            coevaluacion=9,
            autoevaluacion=7,
        )
        self.nota = Nota.objects.create(
            alumno=self.alumno,
            texto='Seguimiento académico privado',
        )
        self.consentimiento = Consentimiento.objects.create(
            alumno=self.alumno,
            tutor_nombre='Tutor del titular',
            parentesco='tutor',
            version_aviso=PRIVACY_NOTICE_VERSION,
            medio='presencial',
        )
        self.datos_solicitud = {
            'tipo': 'acceso',
            'titular_nombre': self.alumno.nombre_completo,
            'numero_cuenta': self.alumno.numero_cuenta,
            'solicitante': 'Solicitante de prueba',
            'parentesco': 'titular',
            'contacto_email': 'solicitante@example.test',
            'descripcion': 'Solicito una copia de mis datos.',
            'confirmacion': 'on',
        }

    def iniciar_sesion_coordinacion(self):
        self.client.force_login(self.user)
        AvisoPrivacidadAceptacion.objects.create(
            usuario=self.user,
            version=PRIVACY_NOTICE_VERSION,
        )

    def crear_solicitud(self, tipo='acceso'):
        return SolicitudARCO.objects.create(
            tipo=tipo,
            titular=self.alumno,
            titular_nombre=self.alumno.nombre_completo,
            numero_cuenta=self.alumno.numero_cuenta,
            solicitante='Solicitante de prueba',
            contacto_email='solicitante@example.test',
            descripcion='Solicitud de prueba.',
            estado='identidad_pendiente',
            fecha_limite=timezone.localdate() + timedelta(days=20),
        )

    def verificar_identidad(self, solicitud):
        return self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'verificar_identidad',
            'titular': self.alumno.id,
            'confirmacion_identidad': 'on',
        })

    def test_public_arco_form_creates_request_and_deadline_without_login(self):
        response = self.client.post(reverse('solicitud_arco'), self.datos_solicitud)

        self.assertEqual(response.status_code, 200)
        solicitud = SolicitudARCO.objects.get()
        self.assertContains(response, f'ARCO-{solicitud.id}')
        self.assertEqual(solicitud.tipo, 'acceso')
        self.assertEqual(solicitud.titular, self.alumno)
        self.assertEqual(solicitud.estado, 'identidad_pendiente')
        self.assertEqual(solicitud.fecha_limite, timezone.localdate() + timedelta(days=20))
        self.assertEqual(solicitud.eventos.count(), 1)

    def test_public_arco_form_does_not_reveal_if_titular_exists(self):
        self.client.post(reverse('solicitud_arco'), self.datos_solicitud)
        self.datos_solicitud['numero_cuenta'] = 'no-existe'

        response = self.client.post(reverse('solicitud_arco'), self.datos_solicitud)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(SolicitudARCO.objects.count(), 2)
        self.assertEqual(
            SolicitudARCO.objects.order_by('-id').first().titular,
            None,
        )
        self.assertContains(response, 'Recibimos tu solicitud')

    def test_arco_coordination_is_limited_to_staff(self):
        regular_user = get_user_model().objects.create_user(
            username='docente-normal',
            password='clave-segura-prueba',
        )
        self.client.force_login(regular_user)
        AvisoPrivacidadAceptacion.objects.create(
            usuario=regular_user,
            version=PRIVACY_NOTICE_VERSION,
        )

        response = self.client.get(reverse('coordinacion_arco'))

        self.assertEqual(response.status_code, 403)

    def test_unverified_identity_cannot_export_student_data(self):
        self.iniciar_sesion_coordinacion()
        solicitud = self.crear_solicitud()

        response = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'exportar_json',
        })

        self.assertRedirects(
            response,
            f"{reverse('coordinacion_arco')}?solicitud={solicitud.id}",
        )
        self.assertEqual(solicitud.eventos.count(), 0)

    def test_identity_verification_uses_already_linked_titular(self):
        self.iniciar_sesion_coordinacion()
        solicitud = self.crear_solicitud()

        page = self.client.get(reverse('coordinacion_arco'), {'solicitud': solicitud.id})
        self.assertContains(page, f'name="titular" value="{self.alumno.id}"')

        response = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'verificar_identidad',
            'confirmacion_identidad': 'on',
        })

        self.assertRedirects(
            response,
            f"{reverse('coordinacion_arco')}?solicitud={solicitud.id}",
        )
        solicitud.refresh_from_db()
        self.assertTrue(solicitud.identidad_verificada)
        self.assertEqual(solicitud.titular, self.alumno)

    def test_identity_verification_without_a_valid_titular_shows_message(self):
        self.iniciar_sesion_coordinacion()
        solicitud = SolicitudARCO.objects.create(
            tipo='acceso',
            titular_nombre='Titular desconocido',
            numero_cuenta='no-existe',
            solicitante='Solicitante de prueba',
            contacto_email='solicitante@example.test',
            descripcion='Solicitud de prueba.',
            estado='identidad_pendiente',
            fecha_limite=timezone.localdate() + timedelta(days=20),
        )

        response = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'verificar_identidad',
            'confirmacion_identidad': 'on',
        }, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Selecciona un titular válido')
        solicitud.refresh_from_db()
        self.assertFalse(solicitud.identidad_verificada)

    def test_verified_access_request_exports_json_and_excel_with_audit(self):
        self.iniciar_sesion_coordinacion()
        solicitud = self.crear_solicitud()
        self.verificar_identidad(solicitud)
        PronosticoLog.objects.create(
            alumno=self.alumno,
            asignatura=self.asignatura,
            version_modelo='tendencia-lineal-v1',
            huella_entradas='e' * 64,
            entradas_agregadas={'parciales': [
                {'parcial': 1, 'promedio': 8},
                {'parcial': 2, 'promedio': 9},
            ]},
            salida={'pronostico_p3': 10},
        )

        response_json = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'exportar_json',
        })
        self.assertEqual(response_json.status_code, 200)
        self.assertEqual(response_json['Content-Type'], 'application/json')
        export_data = json.loads(response_json.content)
        self.assertEqual(export_data['titular']['numero_cuenta'], self.alumno.numero_cuenta)
        self.assertEqual(export_data['calificaciones'][0]['asignatura'], self.asignatura.nombre)
        self.assertEqual(export_data['notas_seguimiento'][0]['texto'], self.nota.texto)
        self.assertEqual(export_data['consentimientos'][0]['tutor'], self.consentimiento.tutor_nombre)
        self.assertEqual(export_data['pronosticos_ia'][0]['salida']['pronostico_p3'], 10)

        response_excel = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'exportar_excel',
        })
        self.assertEqual(response_excel.status_code, 200)
        self.assertIn('spreadsheetml', response_excel['Content-Type'])
        self.assertGreater(len(response_excel.content), 1000)
        workbook = load_workbook(BytesIO(response_excel.content), data_only=True)
        self.assertEqual(workbook['Pronósticos IA']['D2'].value, '{"pronostico_p3": 10}')
        self.assertEqual(
            solicitud.eventos.filter(accion__startswith='exportar_').count(),
            2,
        )

    def test_arco_panel_renders_case_actions_and_export_escapes_excel_formulas(self):
        self.iniciar_sesion_coordinacion()
        solicitud = self.crear_solicitud()
        self.verificar_identidad(solicitud)

        response = self.client.get(
            reverse('coordinacion_arco'),
            {'solicitud': solicitud.id},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Descargar JSON')
        self.assertContains(response, 'Bitácora de actuaciones')

        self.alumno.nombre_completo = '=2+2'
        self.alumno.save(update_fields=['nombre_completo'])
        response = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'exportar_excel',
        })
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook['Titular']['B2'].value, "'=2+2")
        self.assertEqual(workbook['Titular']['B2'].data_type, 's')

    def test_rectification_requires_confirmation_and_logs_field_not_old_value(self):
        self.iniciar_sesion_coordinacion()
        solicitud = self.crear_solicitud('rectificacion')
        self.verificar_identidad(solicitud)

        response = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'rectificar',
            'campo': 'email',
            'valor': 'corregido@example.test',
        })
        self.alumno.refresh_from_db()
        self.assertEqual(self.alumno.email, 'titular@example.test')

        self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'rectificar',
            'campo': 'email',
            'valor': 'corregido@example.test',
            'confirmacion': 'on',
        })
        self.alumno.refresh_from_db()
        self.assertEqual(self.alumno.email, 'corregido@example.test')
        event = solicitud.eventos.get(accion='rectificar')
        self.assertIn('email', event.detalle)
        self.assertNotIn('titular@example.test', event.detalle)

    def test_cancelation_block_prevents_student_data_access(self):
        self.iniciar_sesion_coordinacion()
        solicitud = self.crear_solicitud('cancelacion')
        self.verificar_identidad(solicitud)

        response = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'bloquear',
            'confirmacion': 'BLOQUEAR',
        })

        self.assertRedirects(
            response,
            f"{reverse('coordinacion_arco')}?solicitud={solicitud.id}",
        )
        self.alumno.refresh_from_db()
        self.assertTrue(self.alumno.bloqueado_arco)
        self.assertFalse(alumnos_con_consentimiento().filter(pk=self.alumno.pk).exists())
        self.assertTrue(solicitud.eventos.filter(accion='bloquear').exists())

    def test_suppression_anonymizes_holder_deletes_academic_data_and_keeps_audit(self):
        self.iniciar_sesion_coordinacion()
        solicitud = self.crear_solicitud('cancelacion')
        self.verificar_identidad(solicitud)
        pronostico = PronosticoLog.objects.create(
            alumno=self.alumno,
            asignatura=self.asignatura,
            version_modelo='tendencia-lineal-v1',
            huella_entradas='f' * 64,
            entradas_agregadas={'parciales': []},
            salida={'pronostico_p3': 7},
        )
        revision = SolicitudRevisionPronostico.objects.create(
            alumno=self.alumno,
            pronostico=pronostico,
            solicitante='Tutor',
            parentesco='tutor',
            contacto_email='tutor@example.test',
            motivo='Revisar.',
        )

        response = self.client.post(reverse('coordinacion_arco'), {
            'solicitud': solicitud.id,
            'accion': 'suprimir',
            'confirmacion': 'SUPRIMIR',
        })

        self.assertRedirects(
            response,
            f"{reverse('coordinacion_arco')}?solicitud={solicitud.id}",
        )
        self.alumno.refresh_from_db()
        self.assertTrue(self.alumno.suprimido_arco)
        self.assertTrue(self.alumno.bloqueado_arco)
        self.assertEqual(self.alumno.nombre_completo, f'Registro anonimizado {self.alumno.pk}')
        self.assertEqual(self.alumno.email, '')
        self.assertIsNone(self.alumno.grupo)
        self.assertFalse(Calificacion.objects.filter(alumno=self.alumno).exists())
        self.assertFalse(Nota.objects.filter(alumno=self.alumno).exists())
        self.assertFalse(PronosticoLog.objects.filter(alumno=self.alumno).exists())
        revision.refresh_from_db()
        self.assertIsNone(revision.alumno)
        self.assertIsNone(revision.pronostico)
        self.assertEqual(revision.solicitante, '')
        self.assertEqual(revision.contacto_email, '')
        self.assertEqual(revision.motivo, '')
        self.consentimiento.refresh_from_db()
        self.assertIsNone(self.consentimiento.alumno)
        self.assertEqual(self.consentimiento.tutor_nombre, '')
        solicitud.refresh_from_db()
        self.assertIsNone(solicitud.titular)
        self.assertEqual(solicitud.titular_nombre, '')
        self.assertEqual(solicitud.numero_cuenta, '')
        self.assertEqual(solicitud.solicitante, '')
        self.assertEqual(solicitud.contacto_email, '')
        self.assertEqual(solicitud.estado, 'atendida')
        self.assertTrue(solicitud.eventos.filter(accion='suprimir').exists())
        self.assertNotIn(
            self.alumno.numero_cuenta,
            solicitud.eventos.get(accion='suprimir').detalle,
        )
